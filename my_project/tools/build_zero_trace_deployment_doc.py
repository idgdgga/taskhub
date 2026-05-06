from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "TaskHub_新服务器零痕部署执行手册.docx"

ACCENT = RGBColor(20, 102, 140)
ACCENT_LIGHT = RGBColor(235, 245, 250)
TEXT = RGBColor(34, 40, 49)
MUTED = RGBColor(99, 110, 114)
BORDER = "D9E2EC"
WARNING_FILL = "FFF5E6"
WARNING_BORDER = "F0B458"


def set_east_asia(run, font_name: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_padding(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table) -> None:
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_padding(cell)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 8, "color": BORDER},
                bottom={"val": "single", "sz": 8, "color": BORDER},
                left={"val": "single", "sz": 8, "color": BORDER},
                right={"val": "single", "sz": 8, "color": BORDER},
            )
            if row_index == 0:
                set_cell_shading(cell, "EDF4F9")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = "Arial"
                    set_east_asia(run, "Microsoft YaHei")
                    run.font.color.rgb = TEXT


def add_header_and_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("TaskHub 新服务器零痕部署执行手册")
    run.font.name = "Arial"
    set_east_asia(run, "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(4)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), BORDER)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("第 ")
    fr.font.name = "Arial"
    set_east_asia(fr, "Microsoft YaHei")
    fr.font.size = Pt(9)
    fr.font.color.rgb = MUTED
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    fr._r.append(fld_begin)
    fr._r.append(instr)
    fr._r.append(fld_separate)
    fr._r.append(fld_end)
    fr2 = fp.add_run(" 页")
    fr2.font.name = "Arial"
    set_east_asia(fr2, "Microsoft YaHei")
    fr2.font.size = Pt(9)
    fr2.font.color.rgb = MUTED


def add_title_block(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("TaskHub 新服务器零痕部署执行手册")
    run.font.name = "Arial"
    set_east_asia(run, "Microsoft YaHei")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = ACCENT

    sub = document.add_paragraph()
    sub.paragraph_format.space_after = Pt(16)
    s1 = sub.add_run(
        f"版本：可执行部署版    生成日期：{date.today().isoformat()}    适用场景：宝塔 + Django 后端 + Vue 前端 + Telegram Mini App"
    )
    s1.font.name = "Arial"
    set_east_asia(s1, "Microsoft YaHei")
    s1.font.size = Pt(11)
    s1.font.color.rgb = MUTED

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.style = document.styles["Normal"]
    r = intro.add_run(
        "本手册用于把现有 TaskHub 项目迁移到另一台新服务器，并尽量不留下与当前开发者、旧服务器、旧域名、旧 Bot、旧钱包、旧数据库相关的痕迹。"
        "文档内容基于当前项目代码与配置文件实际结构整理，包含必须替换项、申请入口、提取方式、上传方式、宝塔执行顺序，以及上线后核对清单。"
    )
    r.font.name = "Arial"
    set_east_asia(r, "Microsoft YaHei")
    r.font.size = Pt(11.5)
    r.font.color.rgb = TEXT


def add_heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Arial"
    set_east_asia(run, "Microsoft YaHei")
    run.bold = True
    run.font.color.rgb = TEXT
    run.font.size = Pt(15 if level == 1 else 12.5)


def add_paragraph(document: Document, text: str, *, muted: bool = False) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    run.font.name = "Arial"
    set_east_asia(run, "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED if muted else TEXT


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs:
            run.font.name = "Arial"
            set_east_asia(run, "Microsoft YaHei")
            run.font.size = Pt(11)
            run.font.color.rgb = TEXT
        run = p.add_run(item)
        run.font.name = "Arial"
        set_east_asia(run, "Microsoft YaHei")
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        p = document.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(item)
        run.font.name = "Arial"
        set_east_asia(run, "Microsoft YaHei")
        run.font.size = Pt(11)
        run.font.color.rgb = TEXT


def add_callout(document: Document, title: str, body: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, WARNING_FILL)
    set_cell_padding(cell, top=120, bottom=120, start=160, end=160)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 10, "color": WARNING_BORDER},
        bottom={"val": "single", "sz": 10, "color": WARNING_BORDER},
        left={"val": "single", "sz": 10, "color": WARNING_BORDER},
        right={"val": "single", "sz": 10, "color": WARNING_BORDER},
    )
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(title)
    r1.font.name = "Arial"
    set_east_asia(r1, "Microsoft YaHei")
    r1.bold = True
    r1.font.size = Pt(11.5)
    r1.font.color.rgb = TEXT
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Arial"
    set_east_asia(r2, "Microsoft YaHei")
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = TEXT


def add_code_block(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7FAFC")
    set_cell_padding(cell, top=120, bottom=120, start=140, end=140)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 8, "color": "D9E2EC"},
        bottom={"val": "single", "sz": 8, "color": "D9E2EC"},
        left={"val": "single", "sz": 8, "color": "D9E2EC"},
        right={"val": "single", "sz": 8, "color": "D9E2EC"},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Courier New"
        set_east_asia(run, "Microsoft YaHei")
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT
    document.add_paragraph()


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        p = hdr[idx].paragraphs[0]
        r = p.add_run(header)
        r.font.name = "Arial"
        set_east_asia(r, "Microsoft YaHei")
        r.font.size = Pt(10.5)
        r.bold = True
        r.font.color.rgb = TEXT
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.name = "Arial"
            set_east_asia(r, "Microsoft YaHei")
            r.font.size = Pt(10.2)
            r.font.color.rgb = TEXT
    format_table(table)
    document.add_paragraph()


def build_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    add_header_and_footer(document)
    add_title_block(document)

    add_heading(document, "1. 迁移目标与基本原则")
    add_bullets(
        document,
        [
            "新服务器不保留旧仓库 Git 历史、不保留旧 origin、不保留旧提交作者信息。",
            "新服务器不导入旧数据库，不迁移旧用户、旧钱包、旧充值记录、旧日志、旧上传文件。",
            "Telegram、Apify、TRC20/ERC20/BEP20、客服链接、社群链接、域名、充值钱包全部更换为新资料。",
            "正式上线前，优先保证：能登录后台、能打开 Mini App、能创建任务、能完成任务、能发消息、能自动维护任务状态；充值功能最后再开。",
        ],
    )
    add_callout(
        document,
        "最重要的部署原则",
        "不要在新服务器直接 git clone 当前仓库；最稳的做法是导出一份不带 .git 的源码包，前端只上传新的 dist 静态包，数据库从空库开始。",
    )

    add_heading(document, "2. 必须重新申请、重新生成、重新提取的资料")
    add_paragraph(document, "下面这些资料全部建议重新准备，不要复用旧环境里的任何一个值。")
    add_table(
        document,
        ["项目", "去哪里申请 / 提取", "填到哪里"],
        [
            ["新域名", "域名注册商或现有 DNS 服务商", "后端 ALLOWED_HOSTS / CSRF_TRUSTED_ORIGINS，前端 VITE_API_BASE_URL"],
            ["SSL 证书", "宝塔站点里申请 Let’s Encrypt", "宝塔网站 SSL 配置"],
            ["新数据库", "宝塔 > 数据库 > 新建 MySQL", "后端 .env 或 core/local_settings.py"],
            ["Django SECRET_KEY", "服务器执行 Django 随机密钥生成命令", "后端 .env 的 SECRET_KEY"],
            ["Telegram Bot Token", "Telegram @BotFather 新建机器人", "TELEGRAM_BOT_TOKEN"],
            ["Telegram Bot 用户名", "@BotFather 机器人用户名", "TELEGRAM_BOT_USERNAME"],
            ["Mini App 短名", "@BotFather 的 Main Mini App / Direct Link 设置", "TELEGRAM_MINI_APP_SHORT_NAME"],
            ["Mini App 前端地址", "你的新前端 HTTPS 域名", "TELEGRAM_MINI_APP_URL 与 BotFather Web App URL"],
            ["Webhook Secret", "服务器本地随机生成", "TELEGRAM_WEBHOOK_SECRET"],
            ["Apify Token", "Apify 控制台 API & Integrations", "APIFY_API_TOKEN"],
            ["X 工作号 Cookie", "电脑浏览器登录 X 后提取 auth_token 与 ct0", "APIFY_TWITTER_AUTH_TOKEN / APIFY_TWITTER_CT0"],
            ["TRC20 API Key", "TronGrid Dashboard", "后台充值网络配置 TRC20"],
            ["ERC20 / BEP20 RPC", "自建节点或新的 RPC 服务商", "后台充值网络配置 ERC20/BEP20"],
            ["新地址池助记词", "新建专用钱包", "后台充值网络配置的 HD 主助记词"],
            ["新手续费钱包私钥", "新建专用热钱包", "后台充值网络配置的手续费钱包私钥"],
            ["新归集地址", "新建归集地址或冷钱包地址", "后台充值网络配置的归集目标地址"],
        ],
        [1.2, 2.4, 2.8],
    )

    add_heading(document, "3. 当前项目里必须修改的文件与风险点")
    add_table(
        document,
        ["文件 / 位置", "必须改什么", "原因"],
        [
            ["/my_project/.env 或 core/local_settings.py", "数据库、域名、Telegram、Apify、X、所有生产密钥全部换新", "这是后端运行入口，不能复用旧环境值"],
            ["/my_project/core/settings.py", "通常不直接改代码，但要核对它实际读取了哪些环境变量", "避免漏填必要环境变量"],
            ["/task_vue/.env.production", "改成新 API 域名，或改为同源相对路径", "当前项目这里仍指向旧域名"],
            ["/task_vue/vite.config.ts", "把本地代理目标换成新后端，或改为空白后自行配置", "当前 dev/preview 代理仍指向旧域名"],
            ["/task_vue/.env.development.local", "删除或重建，不要带旧测试 token", "当前文件含真实 bootstrap token，不能带去新环境"],
            ["/announcements/migrations/0005_import_foxigrow_guides.py", "改种子来源、改品牌文案，或在新环境禁用这条种子逻辑", "当前会自动灌入 FoxiGrow 相关内容"],
            ["/announcements/seed/foxigrow_guides.json", "如果保留新手指南种子，要整体替换内容", "避免新库首次迁移后出现旧品牌素材"],
            ["/task_vue/index.html 与 HomeHeader.vue", "改站点标题与前端品牌文案", "当前仍有 TaskFlow 品牌残留"],
            ["后台公告 / 新手指南内容", "群链接、客服、品牌、图片、视频全部重填", "这部分通常存在数据库而不是代码里"],
        ],
        [2.4, 2.0, 2.1],
    )

    add_heading(document, "4. 绝对不要迁移的内容")
    add_bullets(
        document,
        [
            "后端 .git 目录、旧 origin、旧提交历史。",
            "旧服务器 .env、telegram_secrets.py、twitter_secrets.py、apify_secrets.py。",
            "前端 .env.development.local、旧 dist、dist_backup_*、任何构建缓存。",
            "旧数据库：users、wallets、recharge、withdraw、task application、api token、telegram pending、账变、充值地址、审核记录。",
            "旧 media 上传文件、旧日志、旧压缩备份、旧钱包资料。",
        ],
    )
    add_callout(
        document,
        "建议的数据策略",
        "新服务器直接使用空数据库执行 migrate。会员配置、返佣配置、任务模板、公告内容建议人工重建；如果一定要迁配置数据，也要先脱敏并把旧链接、旧域名、旧钱包地址全部改掉。",
    )

    add_heading(document, "5. 宝塔执行版部署步骤")
    add_heading(document, "5.1 本地准备干净发布包", level=2)
    add_numbered(
        document,
        [
            "后端目录复制一份发布版，删除 .git、.env、所有 secrets 文件、logs、media、__pycache__、旧备份包。",
            "前端修改 .env.production、vite.config.ts、品牌文案后执行 npm run build，只保留新的 dist。",
            "如果你后面要重新建 Git 仓库，就在发布包目录里重新 git init，不要沿用当前 origin。",
        ],
    )
    add_heading(document, "5.2 新服务器准备", level=2)
    add_numbered(
        document,
        [
            "在宝塔新建网站，绑定新域名，申请 SSL 证书。",
            "在宝塔新建 MySQL 数据库与专用账号，不要用 root。",
            "准备 Python 运行环境、虚拟环境、Nginx 站点目录、Gunicorn 启动方式。",
        ],
    )
    add_heading(document, "5.3 后端部署", level=2)
    add_numbered(
        document,
        [
            "上传后端干净发布包到新目录，例如 /www/wwwroot/new_taskhub/backend。",
            "在 manage.py 同级创建 .env，或复制 core/local_settings.example.py 为 core/local_settings.py。",
            "填写数据库、SECRET_KEY、ALLOWED_HOSTS、CSRF_TRUSTED_ORIGINS、Telegram、Apify 等环境变量。",
            "执行 pip install -r requirements.txt。",
            "执行 python manage.py migrate。",
            "执行 python manage.py createsuperuser。",
            "执行 python manage.py check，确保无环境缺项。",
        ],
    )
    add_heading(document, "5.4 前端部署", level=2)
    add_numbered(
        document,
        [
            "前端构建前确认 VITE_API_BASE_URL 指向新 API 域名；如果 Nginx 同源反代 /api，也可留空走相对路径。",
            "执行 npm run build。",
            "把 dist 上传到新服务器静态目录，例如 /www/wwwroot/new_taskhub/web/dist。",
            "Nginx 配置前端静态目录并把 /api 反代到 Django/Gunicorn。",
        ],
    )

    add_heading(document, "6. 后端建议填写示例")
    add_paragraph(document, "下面是最常用的一组字段，建议你上线时一条一条核对。")
    add_table(
        document,
        ["变量", "填写示例", "说明"],
        [
            ["MYSQL_DATABASE", "new_taskhub_db", "新建库名"],
            ["MYSQL_USER", "new_taskhub_user", "新建专用用户"],
            ["MYSQL_PASSWORD", "强密码", "不要复用旧库密码"],
            ["SECRET_KEY", "随机 50+ 位字符串", "服务器重新生成"],
            ["DJANGO_DEBUG", "0", "生产环境关闭调试"],
            ["ALLOWED_HOSTS", "new.example.com,127.0.0.1,localhost", "写新域名"],
            ["CSRF_TRUSTED_ORIGINS", "https://new.example.com", "写 HTTPS 来源"],
            ["TELEGRAM_BOT_TOKEN", "BotFather 新机器人 token", "不要复用旧 token"],
            ["TELEGRAM_BOT_USERNAME", "YourNewBot", "不带 @"],
            ["TELEGRAM_MINI_APP_SHORT_NAME", "taskhub_new", "BotFather 的 mini app 短名"],
            ["TELEGRAM_MINI_APP_URL", "https://mini.example.com", "前端 HTTPS 地址"],
            ["TELEGRAM_WEBHOOK_SECRET", "随机字符串", "与 setWebhook 的 secret_token 一致"],
            ["APIFY_API_TOKEN", "Apify 控制台生成", "新的 Apify 账号或新的 token"],
            ["APIFY_TWITTER_AUTH_TOKEN", "从 X 工作号 cookie 提取", "统一平台校验号用"],
            ["APIFY_TWITTER_CT0", "从 X 工作号 cookie 提取", "与 auth_token 同账号同会话"],
        ],
        [2.0, 2.3, 2.0],
    )

    add_heading(document, "7. Telegram、X、充值资料分别怎么获取")
    add_heading(document, "7.1 Telegram Bot / Mini App", level=2)
    add_bullets(
        document,
        [
            "进入 Telegram，搜索 @BotFather。",
            "使用 /newbot 创建一个全新的机器人，获得新的 Bot Token。",
            "在 BotFather 里配置新的头像、名称、用户名、Description、Mini App URL。",
            "如果要配置 Web App / Mini App，使用 BotFather 的对应菜单绑定新前端 HTTPS 地址。",
        ],
    )
    add_paragraph(
        document,
        "官方文档入口：Telegram Bot Features、Bot API setWebhook、Mini Apps 文档。上线时要把 TELEGRAM_WEBHOOK_SECRET 与 setWebhook 的 secret_token 保持一致。",
        muted=True,
    )
    add_heading(document, "7.2 X 的 auth_token 与 ct0", level=2)
    add_bullets(
        document,
        [
            "这两个不是申请来的，而是从你自己的 X 工作号网页登录 Cookie 里提取。",
            "用电脑浏览器登录新的 X 工作号，打开开发者工具，进入 Application 或 Storage。",
            "找到 Cookies 下的 https://x.com，复制 auth_token 与 ct0 两个 Cookie 的 Value。",
            "这两个值必须来自同一个账号、同一次登录会话，不要发给别人。",
        ],
    )
    add_heading(document, "7.3 充值网络资料", level=2)
    add_bullets(
        document,
        [
            "TRC20：去 TronGrid 控制台创建新的 API Key，后台网络配置里填写新的 TRC20 Key。",
            "ERC20 / BEP20：如果不用旧服务商，就重新申请新的 RPC，或者换成你自己的节点。",
            "每条链都重新建一套地址池助记词、一套手续费热钱包、一个新的归集目标地址。",
            "正式环境不要复用测试环境里任何助记词、私钥、归集地址。",
        ],
    )

    add_heading(document, "8. 宝塔计划任务")
    add_paragraph(document, "新环境至少保留下面这些计划任务。充值功能还没开之前，先只挂 maintain_tasks 即可。")
    add_table(
        document,
        ["任务命令", "频率", "作用"],
        [
            ["python manage.py maintain_tasks", "每 1 分钟", "自动失效超时任务、维护虚拟参与人数等"],
            ["python manage.py allocate_recharge_addresses", "每 30 分钟", "批量给用户派生充值地址"],
            ["python manage.py sync_recharge_deposits", "每 1 分钟", "自动扫链确认到账"],
            ["python manage.py sweep_recharge_deposits", "每 1 分钟", "自动归集已到账资金"],
        ],
        [2.9, 1.2, 2.4],
    )

    add_heading(document, "9. 推荐的上线验收顺序")
    add_numbered(
        document,
        [
            "后台能登录，菜单正常打开，SimpleUI 标签不串页。",
            "前端 Mini App 能打开，Telegram 登录正常，语言、任务、排行、个人中心都能显示。",
            "新手指南、客服链接、群链接、品牌名称全部换成新资料。",
            "能创建一个普通任务并完成一次，任务记录状态正常变化。",
            "机器人欢迎消息、群链接按钮、公告频道按钮正常。",
            "最后再开充值：先只开一条链做小额测试，确认自动到账与自动归集都正常，再开其他链。",
        ],
    )

    add_heading(document, "10. 旧服务器收尾清理")
    add_bullets(
        document,
        [
            "停用旧 Bot Token、旧 Apify Token、旧 X Cookie、旧充值钱包。",
            "删除旧服务器 .env、旧日志、旧 media、旧备份压缩包、旧导出的私钥记录。",
            "如果不再保留旧环境，清理 shell history、宝塔计划任务、旧域名解析。",
            "确认新环境稳定后，再考虑关闭旧站点与旧 webhook。",
        ],
    )
    add_callout(
        document,
        "最后提醒",
        "如果你的目标是“和旧身份完全切开”，最关键的是：新域名、新 Bot、新数据库、新充值钱包、新客服链接、新 Git 仓库、新静态包、新品牌内容。只换服务器而不换这些资料，痕迹仍然会保留。",
    )

    add_heading(document, "11. 常用执行命令模板")
    add_paragraph(document, "下面这些命令建议你在新服务器落地时直接参考，按你的实际目录替换路径。")
    add_heading(document, "11.1 生成新的 Django SECRET_KEY", level=2)
    add_code_block(
        document,
        [
            "python -c \"from django.core.management.utils import get_random_secret_key; print(get_random_secret_key())\"",
        ],
    )
    add_heading(document, "11.2 首次后端安装与迁移", level=2)
    add_code_block(
        document,
        [
            "cd /www/wwwroot/new_taskhub/backend",
            "python3 -m venv .venv",
            "source .venv/bin/activate",
            "pip install -r requirements.txt",
            "python manage.py migrate",
            "python manage.py createsuperuser",
            "python manage.py check",
        ],
    )
    add_heading(document, "11.3 设置 Telegram Webhook", level=2)
    add_code_block(
        document,
        [
            "python manage.py telegram_set_webhook --url https://你的域名/api/v1/telegram/webhook/",
            "# 如果命令支持 secret_token，就与 TELEGRAM_WEBHOOK_SECRET 保持一致",
        ],
    )
    add_heading(document, "11.4 前端构建", level=2)
    add_code_block(
        document,
        [
            "cd /path/to/task_vue",
            "npm install",
            "npm run build",
            "# 上传 dist 到新服务器静态目录",
        ],
    )
    add_heading(document, "11.5 自动充值首次准备", level=2)
    add_code_block(
        document,
        [
            "python manage.py diagnose_recharge_network --chain TRC20 --live",
            "python manage.py allocate_recharge_addresses",
            "python manage.py sync_recharge_deposits",
            "python manage.py sweep_recharge_deposits",
        ],
    )
    add_heading(document, "11.6 无 Git 历史发布包建议", level=2)
    add_code_block(
        document,
        [
            "rsync -av --exclude='.git' --exclude='.env' --exclude='media' --exclude='dist_backup_*' 源目录/ 发布目录/",
            "# 或者先复制一份发布目录，再手动删除 .git、.env、logs、cache、旧备份",
        ],
    )

    add_heading(document, "附录：本项目当前已确认的旧环境痕迹")
    add_table(
        document,
        ["位置", "当前状态", "迁移时建议动作"],
        [
            ["后端 Git origin", "仍指向旧 GitHub 仓库", "发布包删除 .git，或新目录 git init 后绑定新远程"],
            ["后端提交作者", "提交历史里带旧作者名与邮箱", "不要把旧仓库历史带到新服务器"],
            ["/task_vue/.env.production", "仍指向旧 API 域名", "改为新域名或改为同源相对路径"],
            ["/task_vue/vite.config.ts", "dev/preview 代理仍指向旧域名", "改到新后端或清空后重配"],
            ["/task_vue/.env.development.local", "含旧测试 token", "删除，不要迁移"],
            ["0005_import_foxigrow_guides.py", "会自动导入 FoxiGrow 教程", "改品牌或禁用旧种子"],
        ],
        [2.5, 1.9, 2.1],
    )
    add_heading(document, "附录：官方入口与控制台")
    add_table(
        document,
        ["服务", "入口", "用途"],
        [
            ["Telegram BotFather", "Telegram 内搜索 @BotFather", "创建新机器人、配置 Mini App"],
            ["Telegram Bot API 文档", "https://core.telegram.org/bots/api", "Webhook、Bot API 参数说明"],
            ["Telegram Mini Apps 文档", "https://core.telegram.org/bots/webapps", "Mini App 登录与 initData 说明"],
            ["Apify 控制台", "https://console.apify.com/settings/integrations", "获取 APIFY_API_TOKEN"],
            ["TronGrid 文档 / 控制台", "https://developers.tron.network/", "申请 TRC20 API Key"],
            ["Geth RPC 文档", "https://geth.ethereum.org/docs/interacting-with-geth/rpc", "自建 ERC20 节点时参考"],
            ["BNB Chain RPC 文档", "https://docs.bnbchain.org/bnb-smart-chain/developers/json_rpc/json-rpc-endpoint/", "BEP20 节点与 RPC 说明"],
            ["Tether Supported Protocols", "https://tether.to/en/supported-protocols/", "确认 USDT 支持链与协议说明"],
        ],
        [1.7, 2.8, 2.0],
    )
    return document


def main() -> None:
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
